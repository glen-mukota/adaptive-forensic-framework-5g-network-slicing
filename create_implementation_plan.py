from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\mukot\OneDrive - North-West University\Documents\Research Proposal Implementation")
DESKTOP = Path(r"C:\Users\mukot\OneDrive - North-West University\Desktop")
SOURCE_RENDER = ROOT / "tmp" / "proposal_pages" / "proposal-1.png"
ASSET_DIR = ROOT / "tmp" / "implementation_plan_assets"
OUTPUT = DESKTOP / "Mukota_Implementation_Action_Plan.docx"

TITLE = "Adaptive Forensic Framework for 5G Network Slicing"
HEADER = TITLE
BLUE = "D9EAF7"
DARK_BLUE = "1F4E79"
LIGHT_BLUE = "EAF3FA"
MID_BLUE = "B9D7EA"
GREY = "F3F5F7"
BLACK = "000000"
WHITE = "FFFFFF"


def font(path, size, bold=False):
    candidates = []
    if bold:
        candidates += [r"C:\Windows\Fonts\timesbd.ttf", r"C:\Windows\Fonts\arialbd.ttf"]
    candidates += [r"C:\Windows\Fonts\times.ttf", r"C:\Windows\Fonts\arial.ttf"]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def crop_logo():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    logo_path = ASSET_DIR / "up_logo_from_proposal.png"
    im = Image.open(SOURCE_RENDER).convert("RGB")
    # Crop the University of Pretoria brand block from the supplied proposal cover.
    crop = im.crop((150, 115, 1010, 395))
    crop.save(logo_path, "PNG")
    return logo_path


def draw_centered(draw, xy, text, fnt, fill="#000000"):
    left, top, right, bottom = xy
    bbox = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=4, align="center")
    x = left + (right - left - (bbox[2] - bbox[0])) / 2
    y = top + (bottom - top - (bbox[3] - bbox[1])) / 2
    draw.multiline_text((x, y), text, font=fnt, fill=fill, spacing=4, align="center")


def draw_architecture():
    target = ASSET_DIR / "implementation_architecture.png"
    image = Image.new("RGB", (1700, 1540), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(None, 34, True)
    layer_font = font(None, 27, True)
    detail_font = font(None, 21)
    draw_centered(draw, (100, 40, 1600, 100), "Implementation Architecture", title_font)

    layers = [
        ("1. Evidence sources", "OAI AMF/SMF/UPF logs | gNodeB/UE events | packet metadata | experiment telemetry"),
        ("2. Acquisition and normalisation", "C# collectors create timestamped, structured evidence records from each approved source"),
        ("3. Slice attribution", "Bind each record to S-NSSAI, slice profile, network function, time and scenario"),
        ("4. Adaptive control", "Explainable rules adjust source priority or capture scope after a defined slice-state event"),
        ("5. Integrity and custody", "SHA-256 at acquisition; custody ledger records source, collector, action, time and storage"),
        ("6. Repository and review", "Evidence index, hash re-verification, metric results and an auditable experiment report"),
    ]
    y = 135
    for index, (heading, detail) in enumerate(layers):
        fill = (234, 243, 250) if index % 2 == 0 else (255, 255, 255)
        draw.rounded_rectangle((120, y, 1580, y + 160), radius=14, fill=fill, outline="black", width=3)
        draw.text((160, y + 24), heading, font=layer_font, fill="black")
        draw.multiline_text((160, y + 78), detail, font=detail_font, fill="black", spacing=4)
        if index < len(layers) - 1:
            midx = 850
            draw.line((midx, y + 160, midx, y + 197), fill="black", width=4)
            draw.polygon([(midx - 10, y + 190), (midx + 10, y + 190), (midx, y + 208)], fill="black")
        y += 225
    image.save(target, "PNG")
    return target


def draw_evidence_lifecycle():
    target = ASSET_DIR / "evidence_lifecycle.png"
    image = Image.new("RGB", (1700, 510), "white")
    draw = ImageDraw.Draw(image)
    title_font = font(None, 32, True)
    label_font = font(None, 22, True)
    detail_font = font(None, 18)
    draw_centered(draw, (80, 25, 1620, 72), "Repeatable Implementation Cycle", title_font)
    steps = [
        ("Prepare", "Configuration\nand slice IDs"),
        ("Generate", "Synthetic traffic\nand scenario ground truth"),
        ("Collect", "Logs, events and\npacket metadata"),
        ("Protect", "Hash and record\ncustody"),
        ("Verify", "Measure, review\nand refine"),
    ]
    x = 85
    for i, (heading, detail) in enumerate(steps):
        draw.rounded_rectangle((x, 145, x + 245, 400), radius=20, fill=(234, 243, 250), outline="black", width=3)
        draw_centered(draw, (x + 10, 175, x + 235, 225), heading, label_font)
        draw_centered(draw, (x + 12, 245, x + 233, 360), detail, detail_font)
        if i < len(steps) - 1:
            start_x = x + 245
            draw.line((start_x + 18, 273, start_x + 80, 273), fill="black", width=4)
            draw.polygon([(start_x + 75, 263), (start_x + 75, 283), (start_x + 93, 273)], fill="black")
        x += 325
    image.save(target, "PNG")
    return target


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        edge_data = kwargs.get(edge)
        tag = qn(f"w:{edge}")
        element = tc_borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        for key in ["val", "sz", "space", "color"]:
            if key in edge_data:
                element.set(qn(f"w:{key}"), str(edge_data[key]))


def set_table_geometry(table, widths_cm):
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total_dxa = sum(int(w * 567) for w in widths_cm)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths_cm):
        grid_col.set(qn("w:w"), str(int(width * 567)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 567)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def remove_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "nil"}, bottom={"val": "nil"},
                left={"val": "nil"}, right={"val": "nil"},
                insideH={"val": "nil"}, insideV={"val": "nil"},
            )


def set_paragraph_keep(paragraph, keep_with_next=False):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_with_next:
        elem = OxmlElement("w:keepNext")
        p_pr.append(elem)
    else:
        elem = OxmlElement("w:keepLines")
        p_pr.append(elem)


def set_run_font(run, name="Times New Roman", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=10)


def add_text(paragraph, text, bold=False, size=10.5, italic=False, color=BLACK):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return run


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    add_text(p, text, size=10.25)
    set_paragraph_keep(p)
    return p


def add_heading(doc, text, level=1):
    style = "Heading 1" if level == 1 else "Heading 2"
    p = doc.add_paragraph(style=style)
    add_text(p, text, bold=True, size=13 if level == 1 else 11.25)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    set_paragraph_keep(p, keep_with_next=True)
    return p


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    add_text(p, text, bold=True, size=10.5)
    set_paragraph_keep(p)
    return p


def add_callout(doc, label, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [16.45])
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_border(cell, top={"val": "single", "sz": "8", "color": "7F9DB9"}, bottom={"val": "single", "sz": "8", "color": "7F9DB9"}, left={"val": "single", "sz": "8", "color": "7F9DB9"}, right={"val": "single", "sz": "8", "color": "7F9DB9"})
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_text(p, label + " ", bold=True, size=10.5)
    add_text(p, body, size=10.5)
    set_paragraph_keep(p)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_step_table(doc, step_no, title, objective, actions, deliverable, acceptance):
    table = doc.add_table(rows=2, cols=3)
    set_table_geometry(table, [1.1, 10.0, 5.35])
    header = table.rows[0].cells
    header[0].text = ""
    header[1].text = ""
    header[2].text = ""
    for cell in header:
        set_cell_shading(cell, BLUE)
        set_cell_border(cell, top={"val": "single", "sz": "8", "color": "7F9DB9"}, bottom={"val": "single", "sz": "8", "color": "7F9DB9"}, left={"val": "single", "sz": "8", "color": "7F9DB9"}, right={"val": "single", "sz": "8", "color": "7F9DB9"})
    p = header[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, str(step_no), bold=True, size=14, color=DARK_BLUE)
    p = header[1].paragraphs[0]
    add_text(p, title, bold=True, size=11.25)
    p = header[2].paragraphs[0]
    add_text(p, "Stage objective", bold=True, size=10.25)
    p = header[2].add_paragraph()
    add_text(p, objective, size=9.5)
    body = table.rows[1].cells
    for cell in body:
        set_cell_border(cell, top={"val": "single", "sz": "6", "color": "A6A6A6"}, bottom={"val": "single", "sz": "6", "color": "A6A6A6"}, left={"val": "single", "sz": "6", "color": "A6A6A6"}, right={"val": "single", "sz": "6", "color": "A6A6A6"})
    p = body[0].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "Action", bold=True, size=9.5, color=DARK_BLUE)
    p = body[1].paragraphs[0]
    add_text(p, actions, size=9.75)
    p = body[2].paragraphs[0]
    add_text(p, "Evidence to show", bold=True, size=9.5, color=DARK_BLUE)
    p = body[2].add_paragraph()
    add_text(p, deliverable, size=9.35)
    p = body[2].add_paragraph()
    add_text(p, "Acceptance: ", bold=True, size=9.25)
    add_text(p, acceptance, size=9.25)
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                set_paragraph_keep(p)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size in [("Heading 1", 13), ("Heading 2", 11.25)]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.paragraph_format.space_after = Pt(0)
    add_text(header_p, HEADER, size=10)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = Pt(0)
    add_page_field(footer_p)
    return doc


def make_document():
    logo = crop_logo()
    architecture = draw_architecture()
    lifecycle = draw_evidence_lifecycle()
    doc = setup_document()

    # Cover page - intentionally follows the supplied proposal's centered identity.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    p.add_run().add_picture(str(logo), width=Inches(5.4))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(32)
    p.paragraph_format.space_after = Pt(8)
    add_text(p, "COS700 Research Proposal", bold=True, size=18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    add_text(p, "Implementation Action Plan", bold=True, size=16)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    add_text(p, TITLE, bold=True, size=17)
    for label, value in [
        ("Student number:", "u21538795"),
        ("Supervisor(s):", "Mr. SM (Sheunesu) Makura"),
        ("Student:", "Lister Glen Mukota"),
        ("Reporting period:", "July - October 2026"),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(7)
        add_text(p, label + " ", bold=True, size=11.5)
        add_text(p, value, size=11.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    add_text(p, "A concise, evidence-led implementation roadmap aligned to the approved proposal.", italic=True, size=10.5)

    doc.add_page_break()

    add_heading(doc, "1 Implementation Direction", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_text(p, "This action plan operationalises the approved Design Science Research artefact. The implementation is a laboratory proof of concept that will collect, attribute, preserve and verify 5G slice evidence without compromising slice isolation. Each stage produces a reviewable artefact so that progress is demonstrated through evidence rather than assertion.", size=10.5)
    add_callout(doc, "Implementation outcome.", "A reproducible C#/.NET prototype integrated with an emulated OpenAirInterface environment. It will ingest selected evidence sources, map records to slice context, adapt collection rules after defined events, hash evidence at acquisition, maintain a custody ledger and report measured results.")

    add_heading(doc, "1.1 Scope Guardrails: What Will and Will Not Be Built", 2)
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [8.2, 8.25])
    hdr = table.rows[0].cells
    for cell, text_value in zip(hdr, ["Included in the proof of concept", "Deliberately outside scope"]):
        set_cell_shading(cell, BLUE)
        set_cell_border(cell, top={"val": "single", "sz": "8", "color": "7F9DB9"}, bottom={"val": "single", "sz": "8", "color": "7F9DB9"}, left={"val": "single", "sz": "8", "color": "7F9DB9"}, right={"val": "single", "sz": "8", "color": "7F9DB9"})
        p = cell.paragraphs[0]
        add_text(p, text_value, bold=True, size=10.5)
    row = table.add_row().cells
    content = [
        "Ubuntu OpenAirInterface laboratory core; AMF, SMF and UPF artefacts; at least two synthetic slices; controlled traffic; explainable rule-based adaptation; SHA-256 integrity verification; chain-of-custody metadata; defined scenarios and metrics.",
        "Live operator deployment; real subscriber data; production-grade TLS interception; unsupervised machine-learning automation; an attempt to reproduce a real-world attack outside the contained laboratory environment.",
    ]
    for cell, text_value in zip(row, content):
        set_cell_border(cell, top={"val": "single", "sz": "6", "color": "A6A6A6"}, bottom={"val": "single", "sz": "6", "color": "A6A6A6"}, left={"val": "single", "sz": "6", "color": "A6A6A6"}, right={"val": "single", "sz": "6", "color": "A6A6A6"})
        p = cell.paragraphs[0]
        add_text(p, text_value, size=9.85)
        p.paragraph_format.line_spacing = 1.03
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    add_heading(doc, "1.2 Target Architecture", 2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_text(p, "The implementation retains the six-layer architecture specified in the proposal, while converting each layer into a demonstrable module and testable output.", size=10.5)
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_after = Pt(1)
    pic_p.add_run().add_picture(str(architecture), width=Cm(13.0))
    add_caption(doc, "Figure 1: Implementation architecture and evidence path")

    doc.add_page_break()
    add_heading(doc, "2 Implementation Method: Six Evidence-Led Steps", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_text(p, "The sequence below controls scope: environment and ground truth come before automation; attribution comes before adaptation; integrity and custody are built into collection rather than added afterward.", size=10.5)
    pic_p = doc.add_paragraph()
    pic_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_p.paragraph_format.space_after = Pt(1)
    pic_p.add_run().add_picture(str(lifecycle), width=Cm(16.0))
    add_caption(doc, "Figure 2: Repeatable implementation cycle used for every controlled scenario")

    add_step_table(
        doc, 1, "Establish a reproducible two-slice laboratory environment",
        "Create the smallest stable OAI environment that can produce labelled evidence.",
        "Configure Ubuntu and OAI core components (AMF, SMF and UPF); define at least two slice profiles; document versions, topology, interfaces, S-NSSAI values and time synchronisation; validate controlled UE/client connectivity.",
        "Version manifest, topology diagram, configuration snapshot and a baseline connectivity log.",
        "Both slice profiles generate labelled synthetic traffic and remain logically separated during baseline checks.",
    )
    add_step_table(
        doc, 2, "Define ground truth and the evidence catalogue",
        "Know which artefacts should exist before collection begins.",
        "Create a scenario ledger that records expected slice events, traffic scripts, evidence locations, timestamps and owner; map OAI logs, session events, control-plane observations, UPF metadata and resource measurements to each scenario.",
        "Expected-artefact register, scenario configuration and time-reference record.",
        "Every experiment has a predetermined artefact list against which collection completeness can be measured.",
    )
    add_step_table(
        doc, 3, "Implement normalised evidence ingestion",
        "Produce structured, reviewable records from approved forensic sources.",
        "Build the C#/.NET collection interface; parse or register source events; retain original source reference; normalise timestamp, source, network function, scenario and preliminary slice fields; store results in an indexed repository.",
        "Prototype build, sample evidence records and a retrieval demonstration.",
        "A baseline run yields queryable records from each selected source without altering the original artefact.",
    )

    doc.add_page_break()
    add_heading(doc, "2 Implementation Method: Six Evidence-Led Steps (continued)", 1)
    add_step_table(
        doc, 4, "Implement slice attribution and transparent rule logic",
        "Associate evidence with the correct slice and explain why the decision was made.",
        "Map records to S-NSSAI, slice type, network function, timestamp and scenario; create explicit rules for missing or ambiguous context; persist the attribution decision and rule identifier with each record.",
        "Attribution report showing inputs, assigned slice context and rule rationale.",
        "Known ground-truth events are assigned to the correct slice, function and scenario with traceable reasoning.",
    )
    add_step_table(
        doc, 5, "Add adaptive control, integrity and chain-of-custody controls",
        "Adapt collection safely while preserving evidential defensibility.",
        "Define observable triggers such as slice creation, load change or security-state change; update collection priority through rule-based logic; calculate SHA-256 at acquisition; record source, collector, action, timestamp, hash and storage location; re-verify on retrieval.",
        "Rule catalogue, custody ledger, hash-verification output and adaptation timing log.",
        "Every collected item has complete custody metadata; stored hashes match on verification; trigger-to-rule-update time is recorded.",
    )
    add_step_table(
        doc, 6, "Execute controlled evaluation and refine the artefact",
        "Assess whether the prototype addresses the proposed forensic challenges under defined conditions.",
        "Run each scenario against documented ground truth; compare expected and collected artefacts; measure attribution, completeness, integrity, adaptation latency, custody completeness and operational overhead; record limitations and refine only where evidence supports change.",
        "Scenario results, metrics table, limitations register and revised prototype/research report input.",
        "Results are reproducible from retained configuration, scripts, ground truth and evidence index.",
    )

    doc.add_page_break()
    add_heading(doc, "3 Controlled Evaluation and Proof of Progress", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_text(p, "The evaluation does not depend on vague demonstrations. Each scenario is run with known inputs and produces measurable evidence aligned to the research questions.", size=10.5)
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [4.4, 7.15, 4.9])
    headers = ["Controlled scenario", "What it demonstrates", "Primary evidence / metric"]
    for cell, text_value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, BLUE)
        set_cell_border(cell, top={"val": "single", "sz": "8", "color": "7F9DB9"}, bottom={"val": "single", "sz": "8", "color": "7F9DB9"}, left={"val": "single", "sz": "8", "color": "7F9DB9"}, right={"val": "single", "sz": "8", "color": "7F9DB9"})
        add_text(cell.paragraphs[0], text_value, bold=True, size=9.8)
    set_repeat_table_header(table.rows[0])
    scenarios = [
        ("Normal slice operation", "Baseline evidence patterns, storage volume, attribution and overhead.", "Completeness; attribution accuracy; CPU/memory/storage."),
        ("Slice allocation manipulation threat model", "Whether a controlled slice-assignment anomaly can be reconstructed from evidence records.", "Cross-source correlation; custody completeness."),
        ("Shared VNF / isolation stress", "Whether collection remains slice-aware without unrelated-slice contamination.", "Correct slice assignment; isolation observations."),
        ("Cross-layer security event", "Correlation of logs, packet metadata and monitoring output across radio/core/virtualisation layers.", "Completeness; evidence timeline."),
        ("Slice creation or load change", "Whether collection rules change promptly after an observable state event.", "Adaptation latency; rule audit trail."),
    ]
    for scenario, demonstration, metric in scenarios:
        cells = table.add_row().cells
        for cell, value in zip(cells, [scenario, demonstration, metric]):
            set_cell_border(cell, top={"val": "single", "sz": "6", "color": "A6A6A6"}, bottom={"val": "single", "sz": "6", "color": "A6A6A6"}, left={"val": "single", "sz": "6", "color": "A6A6A6"}, right={"val": "single", "sz": "6", "color": "A6A6A6"})
            p = cell.paragraphs[0]
            add_text(p, value, size=9.2)
            p.paragraph_format.line_spacing = 1.0
    add_caption(doc, "Table 1: Controlled evaluation scenarios aligned to the approved proposal")

    add_heading(doc, "3.1 Measurement Protocol", 2)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    add_text(p, "All results will be compared with the preconfigured ground truth and retained as part of the experiment record.", size=10.0)
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [4.05, 7.75, 4.65])
    headers = ["Metric", "Operational calculation", "Verification artefact"]
    for cell, text_value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, BLUE)
        set_cell_border(cell, top={"val": "single", "sz": "8", "color": "7F9DB9"}, bottom={"val": "single", "sz": "8", "color": "7F9DB9"}, left={"val": "single", "sz": "8", "color": "7F9DB9"}, right={"val": "single", "sz": "8", "color": "7F9DB9"})
        add_text(cell.paragraphs[0], text_value, bold=True, size=9.25)
    set_repeat_table_header(table.rows[0])
    measures = [
        ("Attribution accuracy", "Correctly assigned records / labelled ground-truth records.", "Attribution report and scenario ledger."),
        ("Collection completeness", "Collected expected artefacts / expected artefacts.", "Expected-artefact register and evidence index."),
        ("Integrity verification", "Stored records whose recomputed SHA-256 matches / hashed records.", "Hash-verification report."),
        ("Adaptation latency", "Elapsed time from defined slice event to matching rule update.", "Timestamped event and rule-change log."),
        ("Operational overhead", "CPU, memory, storage and network delta against baseline.", "Resource-monitoring output."),
        ("Custody completeness", "Records containing all required custody fields / collected records.", "Custody-ledger completeness report."),
    ]
    for metric, calculation, evidence in measures:
        cells = table.add_row().cells
        for cell, value in zip(cells, [metric, calculation, evidence]):
            set_cell_border(cell, top={"val": "single", "sz": "6", "color": "A6A6A6"}, bottom={"val": "single", "sz": "6", "color": "A6A6A6"}, left={"val": "single", "sz": "6", "color": "A6A6A6"}, right={"val": "single", "sz": "6", "color": "A6A6A6"})
            p = cell.paragraphs[0]
            add_text(p, value, size=8.8)
            p.paragraph_format.line_spacing = 1.0
    add_caption(doc, "Table 2: Evaluation measures and the artefacts that will substantiate each result")

    doc.add_page_break()
    add_heading(doc, "4 Milestones, Supervisor Visibility and Achievability", 1)
    add_callout(doc, "Current phase.", "July 2026 is the prototype-build phase in the approved proposal. The milestones below are framed as verifiable gates. They do not assert completion until the listed evidence has been produced and reviewed.")

    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [3.1, 4.45, 5.15, 3.75])
    headers = ["Checkpoint", "Target", "What will be demonstrated", "Supervisor evidence"]
    for cell, text_value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, BLUE)
        set_cell_border(cell, top={"val": "single", "sz": "8", "color": "7F9DB9"}, bottom={"val": "single", "sz": "8", "color": "7F9DB9"}, left={"val": "single", "sz": "8", "color": "7F9DB9"}, right={"val": "single", "sz": "8", "color": "7F9DB9"})
        add_text(cell.paragraphs[0], text_value, bold=True, size=9.5)
    set_repeat_table_header(table.rows[0])
    milestones = [
        ("Environment gate", "30 June 2026 (proposal target)", "OAI core, two slice profiles, controlled client traffic and documented topology.", "Configuration snapshot; version manifest; baseline evidence log."),
        ("Prototype gate", "31 July 2026", "Ingestion, initial attribution, SHA-256 hashing and custody recording.", "Prototype run; sample record; hash and custody verification."),
        ("Evaluation gate", "15 September 2026", "All controlled scenarios completed against ground truth.", "Scenario ledger; metric results; limitations register."),
        ("Report gate", "10 October 2026", "Implementation and evaluation results converted into the research report.", "Draft chapters; figures; evidence-to-result traceability."),
        ("Final readiness", "24-29 October 2026", "Feedback incorporated; final report and poster completed.", "Revision log; final artefact and presentation material."),
    ]
    for row_data in milestones:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_data):
            set_cell_border(cell, top={"val": "single", "sz": "6", "color": "A6A6A6"}, bottom={"val": "single", "sz": "6", "color": "A6A6A6"}, left={"val": "single", "sz": "6", "color": "A6A6A6"}, right={"val": "single", "sz": "6", "color": "A6A6A6"})
            p = cell.paragraphs[0]
            add_text(p, value, size=8.85)
            p.paragraph_format.line_spacing = 1.0
    add_caption(doc, "Table 3: Milestone gates and evidence to be shared with the supervisor")

    add_heading(doc, "4.1 Review Rhythm", 2)
    for bullet in [
        "At each checkpoint, I will submit a one-page evidence update: completed gate, artefacts produced, result summary, issues/risks and the next bounded action.",
        "Configuration, traffic scripts, collection rules and scenario ground truth will be versioned so that an experiment can be repeated and its result traced back to its conditions.",
        "If OAI slicing is unstable, the minimum viable fallback remains the proposal's two-slice eMBB/mMTC configuration; experiments may run sequentially if hardware limits prevent parallel operation.",
        "If OAI logs do not provide sufficient visibility, collection will be expanded to approved raw packet metadata at service-based interfaces, while retaining the same integrity and custody controls.",
    ]:
        add_bullet(doc, bullet)

    add_heading(doc, "5 Statement of Commitment", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_text(p, "The implementation will remain aligned to the approved proposal: a focused, explainable and evaluable proof of concept rather than an overextended production platform. Progress will be shown through configured laboratory artefacts, prototype outputs, custody records, scenario results and explicit limitations. This provides a realistic route from framework design to defensible evaluation.", size=10.5)

    # Ensure all headings and captions stay readable and avoid widows where possible.
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            set_paragraph_keep(paragraph, keep_with_next=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    make_document()
